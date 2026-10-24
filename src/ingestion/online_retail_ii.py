"""Read sale rows from the UCI Online Retail II workbook archive."""

from __future__ import annotations

import hashlib
import shutil
import tempfile
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Iterator
from zipfile import ZipFile
from zoneinfo import ZoneInfo

from openpyxl import load_workbook


DEFAULT_ARCHIVE = Path("data/raw/online+retail+ii.zip")
BUSINESS_TIMEZONE = ZoneInfo("Europe/London")
REQUIRED_COLUMNS = {
    "invoice",
    "stockcode",
    "description",
    "quantity",
    "invoicedate",
    "price",
    "customer id",
    "country",
}


def new_source_stats() -> dict[str, int]:
    return {
        "source_rows": 0,
        "sale_events": 0,
        "cancellations_skipped": 0,
        "returns_or_adjustments_skipped": 0,
        "nonpositive_prices_skipped": 0,
        "invalid_rows_skipped": 0,
        "blank_rows_skipped": 0,
    }


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _number(value) -> Decimal:
    if value is None:
        raise InvalidOperation("missing numeric value")
    return Decimal(str(value).strip())


def _event_timestamp(value: datetime) -> str:
    if not isinstance(value, datetime):
        value = datetime.fromisoformat(str(value).strip())
    if value.tzinfo is None:
        value = value.replace(tzinfo=BUSINESS_TIMEZONE)
    else:
        value = value.astimezone(BUSINESS_TIMEZONE)
    return value.isoformat()


def iter_online_retail_sales(
    archive_path: str | Path = DEFAULT_ARCHIVE,
    stats: dict[str, int] | None = None,
) -> Iterator[dict]:
    """Yield valid positive sale lines; keep the source workbook on disk.

    Cancellation, return, and nonpositive-price rows are counted and skipped
    because the existing real-time forecast/anomaly pipeline models sales.
    The source archive should remain unmodified so those rows can be analyzed
    separately when the pipeline is extended for refunds and cancellations.
    """
    archive_path = Path(archive_path)
    if not archive_path.is_file():
        raise FileNotFoundError(f"Online Retail II archive not found: {archive_path}")

    counters = stats if stats is not None else new_source_stats()

    with ZipFile(archive_path) as archive:
        workbook_members = [
            name for name in archive.namelist()
            if name.lower().endswith(".xlsx") and not name.startswith("__MACOSX/")
        ]
        if len(workbook_members) != 1:
            raise ValueError(
                "Expected one .xlsx workbook inside the Online Retail II ZIP; "
                f"found {len(workbook_members)}."
            )

        with tempfile.TemporaryDirectory(prefix="online-retail-ii-") as temp_dir:
            workbook_path = Path(temp_dir) / Path(workbook_members[0]).name
            with archive.open(workbook_members[0]) as source_file:
                with workbook_path.open("wb") as destination_file:
                    shutil.copyfileobj(source_file, destination_file)

            workbook = load_workbook(
                workbook_path,
                read_only=True,
                data_only=True,
            )
            try:
                for sheet in workbook.worksheets:
                    rows = sheet.iter_rows(values_only=True)
                    header = next(rows, None)
                    if not header:
                        continue

                    normalized_header = [
                        _text(value).lower() for value in header
                    ]
                    missing = REQUIRED_COLUMNS - set(normalized_header)
                    if missing:
                        raise ValueError(
                            f"Worksheet {sheet.title!r} is missing columns: "
                            f"{sorted(missing)}"
                        )
                    positions = {
                        name: normalized_header.index(name)
                        for name in REQUIRED_COLUMNS
                    }

                    for excel_row, row in enumerate(rows, start=2):
                        if not row or all(value is None for value in row):
                            counters["blank_rows_skipped"] += 1
                            continue
                        counters["source_rows"] += 1

                        try:
                            invoice_id = _text(row[positions["invoice"]])
                            stock_code = _text(row[positions["stockcode"]])
                            quantity_decimal = _number(row[positions["quantity"]])
                            unit_price_decimal = _number(row[positions["price"]])
                            invoice_date = row[positions["invoicedate"]]

                            if not invoice_id or not stock_code:
                                counters["invalid_rows_skipped"] += 1
                                continue
                            if invoice_id.upper().startswith("C"):
                                counters["cancellations_skipped"] += 1
                                continue
                            if quantity_decimal <= 0:
                                counters["returns_or_adjustments_skipped"] += 1
                                continue
                            if unit_price_decimal <= 0:
                                counters["nonpositive_prices_skipped"] += 1
                                continue
                            if quantity_decimal != quantity_decimal.to_integral_value():
                                counters["invalid_rows_skipped"] += 1
                                continue

                            quantity = int(quantity_decimal)
                            revenue = (quantity_decimal * unit_price_decimal).quantize(
                                Decimal("0.01"), rounding=ROUND_HALF_UP
                            )
                            country = _text(row[positions["country"]]) or "Unknown"
                            customer = _text(row[positions["customer id"]]) or None
                            description = _text(row[positions["description"]]) or None
                            timestamp = _event_timestamp(invoice_date)
                            business_date = (
                                invoice_date.date().isoformat()
                                if isinstance(invoice_date, datetime)
                                else datetime.fromisoformat(str(invoice_date)).date().isoformat()
                            )
                            row_key = (
                                f"{sheet.title}|{excel_row}|{invoice_id}|{stock_code}"
                            )
                            event_id = "online_retail_ii_" + hashlib.sha256(
                                row_key.encode("utf-8")
                            ).hexdigest()[:24]

                            counters["sale_events"] += 1
                            yield {
                                "event_id": event_id,
                                "timestamp": timestamp,
                                "business_date": business_date,
                                "invoice_id": invoice_id,
                                "source_row": excel_row,
                                "source_sheet": sheet.title,
                                "source_dataset": "Online Retail II",
                                "customer_id": customer,
                                "product_id": stock_code,
                                "description": description,
                                "category": "Unclassified",
                                # Legacy pipeline field: this is Country, not a store.
                                "store_id": country,
                                "country": country,
                                "region": country,
                                "quantity": quantity,
                                "unit_price": float(unit_price_decimal),
                                # The source does not provide discount or payment data.
                                "discount": None,
                                "payment_method": None,
                                "channel": "online",
                                "revenue": float(revenue),
                            }
                        except (IndexError, InvalidOperation, TypeError, ValueError):
                            counters["invalid_rows_skipped"] += 1
                            continue
            finally:
                workbook.close()
